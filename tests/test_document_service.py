from io import BytesIO

from docx import Document

from app.services.document import collect_document_paragraphs, translate_document


def test_translate_document_preserves_paragraph_count_and_replaces_text():
    doc = Document()
    doc.add_paragraph("Hello world")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "In table"

    source = BytesIO()
    doc.save(source)

    refs = collect_document_paragraphs(Document(BytesIO(source.getvalue())))
    assert [r.source_text for r in refs] == ["Hello world", "In table"]

    translated_bytes = translate_document(source.getvalue(), ["Bonjour", "Dans le tableau"])

    translated = Document(BytesIO(translated_bytes))
    assert translated.paragraphs[0].text == "Bonjour"
    assert translated.tables[0].cell(0, 0).text == "Dans le tableau"
