from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph


@dataclass
class ParagraphRef:
    paragraph: Paragraph
    source_text: str


def _iter_container_paragraphs(container: DocumentType | _Cell) -> Iterable[Paragraph]:
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from _iter_container_paragraphs(cell)


def collect_document_paragraphs(document: DocumentType) -> list[ParagraphRef]:
    refs: list[ParagraphRef] = []
    for paragraph in _iter_container_paragraphs(document):
        text = paragraph.text.strip()
        if text:
            refs.append(ParagraphRef(paragraph=paragraph, source_text=text))

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                refs.append(ParagraphRef(paragraph=paragraph, source_text=paragraph.text.strip()))
        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                refs.append(ParagraphRef(paragraph=paragraph, source_text=paragraph.text.strip()))
    return refs


def replace_paragraph_text(paragraph: Paragraph, translated_text: str) -> None:
    if not paragraph.runs:
        paragraph.text = translated_text
        return

    paragraph.runs[0].text = translated_text
    for run in paragraph.runs[1:]:
        run.text = ""


def translate_document(file_bytes: bytes, translated_texts: list[str]) -> bytes:
    document = Document(BytesIO(file_bytes))
    refs = collect_document_paragraphs(document)

    if len(refs) != len(translated_texts):
        raise ValueError("Translation count mismatch while rebuilding the document.")

    for ref, translated in zip(refs, translated_texts, strict=True):
        replace_paragraph_text(ref.paragraph, translated)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
